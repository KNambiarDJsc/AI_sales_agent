"""Client console API (apps/api/routers/console.py) against the real Postgres, with a
fake telephony provider — nothing is dialled. Each test gets its own console tenant."""
import io
import uuid

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from apps.api.main import app
from config.settings import get_settings
from telephony.base import CallStatus, OutboundCallResult
from tests.unit.test_frejun_routes import _drop_pool
from tests.unit.test_media_exotel import requires_db

PASSWORD = "test-console-pass"


class FakeTelephony:
    name = "fake"

    def __init__(self):
        self.calls, self.hung_up = [], []
        self.live = CallStatus.IN_PROGRESS

    async def create_outbound_call(self, request):
        self.calls.append(request)
        return OutboundCallResult(provider_call_id=f"fake-{uuid.uuid4().hex[:8]}", status=CallStatus.QUEUED)

    async def hangup_call(self, provider_call_id):
        self.hung_up.append(provider_call_id)

    async def get_call_status(self, provider_call_id):
        return self.live


@pytest.fixture
def console(monkeypatch):
    _drop_pool()
    s = get_settings()
    monkeypatch.setattr(s, "console_password", PASSWORD)
    monkeypatch.setattr(s, "console_username", "admin")
    monkeypatch.setattr(s, "console_tenant_name", f"Console test {uuid.uuid4().hex[:8]}")
    monkeypatch.setattr(s, "public_base_url", "https://agent.example.com")
    telephony = FakeTelephony()
    monkeypatch.setattr("apps.api.routers.console.get_telephony_provider", lambda name=None: telephony)
    with TestClient(app) as client:
        token = client.post("/console/login", json={"username": "admin", "password": PASSWORD}).json()["token"]
        client.headers["Authorization"] = f"Bearer {token}"
        yield client, telephony
    _drop_pool()


@requires_db
def test_login_is_required_and_a_wrong_password_is_refused(monkeypatch):
    _drop_pool()
    monkeypatch.setattr(get_settings(), "console_password", PASSWORD)
    monkeypatch.setattr(get_settings(), "console_username", "admin")
    with TestClient(app) as client:
        assert client.get("/console/overview").status_code == 401
        assert client.post("/console/login", json={"username": "admin", "password": "nope"}).status_code == 401
        assert client.post("/console/login", json={"username": "someone", "password": PASSWORD}).status_code == 401
        assert client.post("/console/login", json={"username": "Admin", "password": PASSWORD}).status_code == 200
        assert client.get("/console/overview", headers={"Authorization": "Bearer forged.token"}).status_code == 401
    _drop_pool()


@requires_db
def test_add_list_and_remove_a_lead(console):
    client, _ = console
    added = client.post("/console/leads", json={"name": "Naman", "phone": "98765 43210", "business": "Demo Business"})
    assert added.status_code == 200
    leads = client.get("/console/leads").json()
    assert [(l["name"], l["phone"], l["business"], l["result"]) for l in leads] == [
        ("Naman", "+919876543210", "Demo Business", "Not called yet")]
    assert client.post("/console/leads", json={"phone": "12"}).status_code == 422  # invalid number
    assert client.delete(f"/console/leads/{added.json()['id']}").json() == {"deleted": True}
    assert client.get("/console/leads").json() == []


@requires_db
def test_excel_upload_imports_valid_rows_and_reports_bad_ones(console):
    client, _ = console
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Phone", "Business", "Category"])
    ws.append(["A", "+91 98111 11111", "Shop A", "Sarees"])
    ws.append(["B", "not a number", "Shop B", None])
    ws.append(["C", "98222 22222", "Shop C", None])
    buf = io.BytesIO()
    wb.save(buf)
    resp = client.post("/console/leads/upload",
                       files={"file": ("leads.xlsx", buf.getvalue(), "application/vnd.ms-excel")})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["created"] == 2 and len(body["errors"]) == 1
    assert {(l["name"], l["category"]) for l in client.get("/console/leads").json()} == {("A", "Sarees"), ("C", "")}


@requires_db
def test_test_call_dials_once_and_shows_up_as_live(console):
    client, telephony = console
    resp = client.post("/console/calls/test", json={"phone": "9876543210", "name": "Naman", "business": "Demo Business"})
    assert resp.status_code == 200, resp.text
    assert len(telephony.calls) == 1 and telephony.calls[0].to_number == "+919876543210"
    live = client.get("/console/overview").json()["live"]
    assert live["name"] == "Naman" and live["test"] is True
    assert client.post("/console/calls/test", json={"phone": "9876543210"}).status_code == 409  # one at a time


@requires_db
def test_overview_template_and_export_work(console):
    client, _ = console
    client.post("/console/leads", json={"name": "Naman", "phone": "9876543211", "business": "Demo Business"})
    overview = client.get("/console/overview").json()
    assert overview["leads"]["pending"] == 1 and overview["campaign"]["status"] == "paused"
    template = load_workbook(io.BytesIO(client.get("/console/leads/template.xlsx").content))
    assert [c.value for c in template.active[1]] == ["Name", "Phone", "Business", "Product category"]
    export = load_workbook(io.BytesIO(client.get("/console/export.xlsx").content))
    assert export.sheetnames[0] == "All leads"
    assert [c.value for c in export["All leads"][2]][:4] == ["Naman", "Demo Business", "+919876543211", "Not called yet"]
    assert client.post("/console/campaign/start").json() == {"status": "active"}
    assert client.put("/console/campaign/hours", json={"start": "10:00", "end": "18:30"}).status_code == 200
    assert client.get("/console/overview").json()["campaign"]["call_window_end"] == "18:30"


@requires_db
def test_end_call_button_hangs_up_and_the_call_stops_showing_live(console):
    client, telephony = console
    attempt_id = client.post("/console/calls/test", json={"phone": "9876543212"}).json()["attempt_id"]
    assert client.post(f"/console/calls/{attempt_id}/hangup").json() == {"ended": True}
    assert len(telephony.hung_up) == 1
    assert client.get("/console/overview").json()["live"] is None  # no webhook needed
    assert client.get(f"/console/calls/{attempt_id}").json()["status"] == "canceled"
    assert client.post("/console/calls/test", json={"phone": "9876543212"}).status_code == 200  # not blocked


@requires_db
def test_a_call_the_provider_says_is_over_stops_showing_live(console):
    # Vobiz once kept a call open after the customer hung up and sent no webhook.
    from datetime import timedelta

    from apps.api.routers import console as console_module
    from database.models import CallAttempt
    from database.session import session_scope

    client, telephony = console
    attempt_id = client.post("/console/calls/test", json={"phone": "9876543213"}).json()["attempt_id"]

    async def age_and_reconcile(times):
        async with session_scope() as session:
            attempt = await session.get(CallAttempt, uuid.UUID(attempt_id))
            attempt.created_at = attempt.created_at - timedelta(seconds=90)
        for _ in range(times):
            await console_module._reconcile_live_calls()

    client.portal.call(age_and_reconcile, 2)  # provider says still live: untouched
    assert client.get("/console/overview").json()["live"]["attempt_id"] == attempt_id
    telephony.live = CallStatus.UNKNOWN
    client.portal.call(age_and_reconcile, 1)  # one "not live" answer isn't enough
    assert client.get("/console/overview").json()["live"] is not None
    client.portal.call(age_and_reconcile, 1)
    assert client.get("/console/overview").json()["live"] is None
    assert client.get(f"/console/calls/{attempt_id}").json()["status"] == "no_answer"
