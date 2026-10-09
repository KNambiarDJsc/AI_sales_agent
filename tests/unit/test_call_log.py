"""Excel call log (workers/call_log.py): the yes/no decision and the workbook layout."""
from datetime import datetime
from types import SimpleNamespace

from openpyxl import load_workbook

from workers.call_log import CallRow, _agreed, write_workbook


def _attempt(outcome=None, status="completed"):
    return SimpleNamespace(outcome=outcome, status=status)


def test_agreed_comes_from_the_state_path_not_from_the_llm():
    followup = {"INTERESTED"}
    path = ["INTRO", "PERMISSION", "DISCOVERY", "QUALIFICATION", "INTERESTED", "END"]
    assert _agreed(_attempt(), path, followup, None, None) == "Yes"
    assert _agreed(_attempt(), ["INTRO", "PERMISSION", "NOT_INTERESTED", "END"], followup, None, None) == "No"
    assert _agreed(_attempt(), ["INTRO", "UNCERTAIN"], followup, SimpleNamespace(sales_followup_required=True), None) == "Yes"
    assert _agreed(_attempt(), ["INTRO", "CALLBACK"], followup, None, object()) == "Callback requested"
    assert _agreed(_attempt(outcome="do_not_call"), path, followup, None, None) == "Do not call"  # DNC always wins
    assert _agreed(_attempt(status="no_answer"), [], followup, None, None) == "No answer"


def test_demo_test_and_simulated_calls_are_left_out_of_the_log():
    from workers.call_log import _is_logged

    assert _is_logged("vobiz", "Vobiz Test Campaign")  # the real test call
    assert _is_logged("vobiz", "Dio Perfumes outreach")
    assert not _is_logged("browser_demo", "Browser Demo Campaign")
    assert not _is_logged("vobiz", "Vobiz Simulation Campaign")
    assert not _is_logged("frejun", "frejun-test")
    assert not _is_logged("fake", "anything")


def _row(agreed, name="Naman"):
    return CallRow(
        call_time=datetime(2026, 10, 9, 11, 30), name=name, business="Dio Perfumes", phone="+919800000000",
        agreed=agreed, outcome="completed", call_status="completed", talk_seconds=95, final_state="END",
        path="INTRO → INTERESTED → END", callback_time="", summary="", facts="",
        transcript="Agent: Hello\nCustomer: Yes please call me", last_customer_words="Yes please call me",
        campaign="Test", provider="vobiz", attempt_id="a1",
    )


def test_workbook_has_every_call_and_a_sheet_of_the_yeses(tmp_path):
    path = tmp_path / "call_log.xlsx"
    write_workbook([_row("Yes", "Naman"), _row("No", "Other")], path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Calls", "Agreed to sales team"]
    calls = list(wb["Calls"].iter_rows(values_only=True))
    assert calls[0][:5] == ("Call time", "Name", "Business", "Phone", "Agreed to sales team?")
    assert [r[1] for r in calls[1:]] == ["Other", "Naman"]  # newest first
    yes = list(wb["Agreed to sales team"].iter_rows(values_only=True))
    assert len(yes) == 2 and yes[1][1] == "Naman" and yes[1][4] == "Yes please call me"
