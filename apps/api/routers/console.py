"""Client console API — what the web/ Next.js app calls.

Everything here needs a login (apps/api/console_auth.py). The console works on two
campaigns of its own, created on first use under a "Console" tenant:

- **Leads** — the client's uploaded / added leads. While it is "running", the
  in-process dialer (`run_console_dialer`) calls them one at a time within its calling
  hours, using the same `dispatch_next_call` the standalone worker uses (DNC
  suppression, idempotent attempts, lead lifecycle all unchanged).
- **Test calls** — "call this number now": one call, any time of day, outside the
  queue.

Results come from the same rules as the Excel call log (workers/call_log.py):
Interested / Callback requested / Not interested / Do not call / Unclear / No answer.
"""
from __future__ import annotations

import asyncio
import contextlib
import io
import logging
import time
import uuid
from datetime import datetime, timezone

import pandas as pd
import phonenumbers
from fastapi import APIRouter, Depends, HTTPException, Response, UploadFile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel
from sqlalchemy import func, select

from apps.api.console_auth import issue_token, password_is_correct, require_console
from config.settings import get_settings
from database.models import CallAttempt, Campaign, Conversation, Lead, Tenant, TranscriptSegment, Turn
from database.session import session_scope
from database.repositories.call_repository import CallRepository
from services.call_worker.dispatch import dispatch_next_call
from services.call_worker.lifecycle import finalize_call
from services.campaign.service import CampaignService
from services.lead_import.importer import guess_column_mapping, import_leads
from telephony.base import CallStatus
from telephony.factory import get_telephony_provider
from workers.call_log import CallRow, collect_rows, write_workbook

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/console", tags=["console"])
public_router = APIRouter(prefix="/console", tags=["console"])

TENANT_NAME = "Console"
LEADS_CAMPAIGN = "Leads"
TEST_CAMPAIGN = "Test calls"
ACTIVE_CALL_STATUSES = ("queued", "dialing", "ringing", "in_progress")
DIALER_POLL_SECONDS = 4
# Every this many dialer ticks, check calls we think are live against the provider.
RECONCILE_EVERY_TICKS = 5
# How long a call may look live before we start asking the provider whether it really is.
RECONCILE_AFTER_SECONDS = 45


# --------------------------------------------------------------------------- setup


async def _campaigns(session) -> tuple[Campaign, Campaign]:
    """(leads campaign, test campaign), created on first use."""
    s = get_settings()
    tenant_name = s.console_tenant_name or TENANT_NAME
    tenant = (await session.execute(select(Tenant).where(Tenant.name == tenant_name))).scalar_one_or_none()
    if tenant is None:
        tenant = Tenant(name=tenant_name)
        session.add(tenant)
        await session.flush()
    found = {
        c.name: c
        for c in (await session.execute(select(Campaign).where(Campaign.tenant_id == tenant.id))).scalars()
    }
    for name, window in ((LEADS_CAMPAIGN, (s.default_call_window_start, s.default_call_window_end)),
                         (TEST_CAMPAIGN, ("00:00", "23:59"))):
        if name not in found:
            found[name] = Campaign(
                tenant_id=tenant.id, name=name, script_id=s.console_script_id, script_version=1,
                status="paused", telephony_provider=s.telephony_provider,
                call_window_start=window[0], call_window_end=window[1], timezone=s.default_timezone,
                max_concurrent_calls=1, max_call_duration_seconds=s.default_max_call_duration_seconds,
                config={"console": True, "campaign_prompt_path": s.console_campaign_prompt},
            )
            session.add(found[name])
        if found[name].max_call_duration_seconds != s.default_max_call_duration_seconds:
            found[name].max_call_duration_seconds = s.default_max_call_duration_seconds
        if found[name].script_id != s.console_script_id:  # the configured flow changed
            found[name].script_id = s.console_script_id
            found[name].config = {**found[name].config, "campaign_prompt_path": s.console_campaign_prompt}
    await session.flush()
    return found[LEADS_CAMPAIGN], found[TEST_CAMPAIGN]


def _public_base() -> str:
    return get_settings().effective_public_base_url


def _normalize_phone(raw: str) -> str:
    try:
        parsed = phonenumbers.parse(raw, "IN")
    except phonenumbers.NumberParseException:
        parsed = None
    if parsed is None or not phonenumbers.is_valid_number(parsed):
        raise HTTPException(status_code=422, detail=f"Not a valid phone number: {raw!r}")
    return phonenumbers.format_number(parsed, phonenumbers.PhoneNumberFormat.E164)


async def _dispatch_one(session, campaign: Campaign) -> bool:
    base = _public_base()
    if not base.startswith("https://"):
        raise HTTPException(status_code=503, detail="Server has no public https URL (PUBLIC_BASE_URL) — start the tunnel")
    return await dispatch_next_call(
        session, campaign, get_telephony_provider(campaign.telephony_provider),
        media_websocket_base_url=f"{base.replace('https://', 'wss://')}/media",
        status_callback_base_url=f"{base}/webhooks",
    )


# --------------------------------------------------------------------------- auth


class LoginBody(BaseModel):
    username: str = ""
    password: str


@public_router.post("/login")
async def login(body: LoginBody) -> dict:
    if not get_settings().console_password:
        raise HTTPException(status_code=503, detail="Console disabled: set CONSOLE_PASSWORD in .env")
    if not password_is_correct(body.password, body.username):
        await asyncio.sleep(1.0)  # slow down guessing
        raise HTTPException(status_code=401, detail="Wrong username or password")
    return {"token": issue_token(), "expires_in_hours": get_settings().console_token_hours}


# --------------------------------------------------------------------------- overview


def _row_dict(row: CallRow) -> dict:
    return {
        "attempt_id": row.attempt_id, "lead_id": row.lead_id, "time": row.call_time.isoformat() if row.call_time else None,
        "name": row.name, "business": row.business, "phone": row.phone, "result": row.agreed,
        "outcome": row.outcome, "status": row.call_status, "talk_seconds": row.talk_seconds,
        "final_state": row.final_state, "path": row.path, "callback_time": row.callback_time,
        "summary": row.summary, "campaign": row.campaign,
    }


async def _console_rows() -> list[CallRow]:
    async with session_scope() as session:
        leads_c, test_c = await _campaigns(session)
        ids = {leads_c.id, test_c.id}
    return await collect_rows(campaign_ids=ids, apply_skip_filter=False)


async def _live_call() -> dict | None:
    """The call in progress (any console campaign), with its latest transcript lines."""
    async with session_scope() as session:
        leads_c, test_c = await _campaigns(session)
        attempt = (await session.execute(
            select(CallAttempt).where(CallAttempt.campaign_id.in_([leads_c.id, test_c.id]),
                                      CallAttempt.status.in_(ACTIVE_CALL_STATUSES))
            .order_by(CallAttempt.created_at.desc()).limit(1)
        )).scalar_one_or_none()
        if attempt is None:
            return None
        age = (datetime.now(timezone.utc) - attempt.created_at).total_seconds()
        if age > get_settings().default_max_call_duration_seconds + 180:
            return None  # stale row from a call whose end was never reported
        lead = await session.get(Lead, attempt.lead_id)
        conv = (await session.execute(
            select(Conversation).where(Conversation.call_attempt_id == attempt.id)
            .order_by(Conversation.created_at.desc()).limit(1)
        )).scalar_one_or_none()
        lines = []
        if conv is not None:
            turns = (await session.execute(
                select(Turn).where(Turn.conversation_id == conv.id).order_by(Turn.turn_index.desc()).limit(6)
            )).scalars().all()
            texts = {s.turn_id: s.text for s in (await session.execute(
                select(TranscriptSegment).where(TranscriptSegment.turn_id.in_([t.id for t in turns]))
            )).scalars()}
            lines = [{"speaker": t.speaker, "text": texts.get(t.id, "")} for t in reversed(turns) if texts.get(t.id)]
        return {
            "attempt_id": str(attempt.id), "name": lead.contact_name if lead else "", "phone": lead.phone_e164 if lead else "",
            "business": lead.business_name if lead else "", "status": "talking" if conv is not None else "ringing",
            "seconds": int(age), "state": conv.current_state if conv else None, "lines": lines,
            "test": attempt.campaign_id == test_c.id,
        }


@router.get("/overview", dependencies=[Depends(require_console)])
async def overview() -> dict:
    rows = await _console_rows()
    async with session_scope() as session:
        leads_c, _ = await _campaigns(session)
        by_status = dict((await session.execute(
            select(Lead.status, func.count()).where(Lead.campaign_id == leads_c.id).group_by(Lead.status)
        )).all())
        campaign = {"status": leads_c.status, "call_window_start": leads_c.call_window_start,
                    "call_window_end": leads_c.call_window_end, "timezone": leads_c.timezone}
    results: dict[str, int] = {}
    for row in rows:
        results[row.agreed] = results.get(row.agreed, 0) + 1
    today = datetime.now().date()
    return {
        "campaign": campaign,
        "within_calling_hours": _within_window(campaign),
        "max_call_seconds": get_settings().default_max_call_duration_seconds,
        "leads": {"total": sum(by_status.values()), "pending": by_status.get("pending", 0) + by_status.get("queued", 0),
                  "by_status": by_status},
        "results": results,
        "calls_total": len(rows),
        "calls_today": sum(1 for r in rows if r.call_time and r.call_time.date() == today),
        "talk_seconds": sum(r.talk_seconds or 0 for r in rows),
        "live": await _live_call(),
        "recent": [_row_dict(r) for r in reversed(rows[-8:])],
    }


def _within_window(campaign: dict) -> bool:
    from zoneinfo import ZoneInfo

    now = datetime.now(ZoneInfo(campaign["timezone"])).strftime("%H:%M")
    return campaign["call_window_start"] <= now <= campaign["call_window_end"]


# --------------------------------------------------------------------------- leads


@router.get("/leads", dependencies=[Depends(require_console)])
async def list_leads() -> list[dict]:
    rows = await _console_rows()
    latest: dict[str, CallRow] = {}
    for row in rows:
        latest[row.lead_id] = row
    async with session_scope() as session:
        leads_c, _ = await _campaigns(session)
        leads = (await session.execute(
            select(Lead).where(Lead.campaign_id == leads_c.id).order_by(Lead.created_at.desc())
        )).scalars().all()
        out = []
        for lead in leads:
            last = latest.get(str(lead.id))
            result = last.agreed if last else ("Do not call" if lead.status == "suppressed" else
                                               "Calling…" if lead.status in ("queued", "in_progress") else "Not called yet")
            out.append({
                "id": str(lead.id), "name": lead.contact_name or "", "business": lead.business_name or "",
                "phone": lead.phone_e164, "status": lead.status, "attempts": lead.attempts_count, "result": result,
                "last_call_id": last.attempt_id if last else None, "summary": last.summary if last else "",
                "added": lead.created_at.isoformat(), "category": (lead.extra or {}).get("product_category", ""),
            })
        return out


class NewLead(BaseModel):
    name: str = ""
    phone: str
    business: str = ""
    category: str = ""  # product category: personalizes the pitch (lead_fields.product_category)


# Lead-sheet headers that mean "product category"; stored as Lead.extra["product_category"],
# the name the script's prompts use. Other unrecognized columns stay in extra as-is.
_CATEGORY_HEADERS = {"product category", "category", "product", "products", "product type", "what they sell"}


def _with_category(extra: dict, category: str) -> dict:
    return {**extra, "product_category": category.strip()} if category.strip() else extra


@router.post("/leads", dependencies=[Depends(require_console)])
async def add_lead(body: NewLead) -> dict:
    phone = _normalize_phone(body.phone)
    async with session_scope() as session:
        leads_c, _ = await _campaigns(session)
        service = CampaignService(session)
        if await service.leads.is_suppressed(leads_c.tenant_id, phone):
            raise HTTPException(status_code=409, detail="This number asked not to be called (do-not-call list)")
        existing = await service.leads.find_by_dedupe_key(leads_c.id, phone)
        if existing is not None:
            existing.contact_name = body.name or existing.contact_name
            existing.business_name = body.business or existing.business_name
            existing.extra = _with_category(existing.extra, body.category)
            if existing.status not in ("queued", "in_progress", "suppressed"):
                existing.status = "pending"  # back in the queue
            return {"id": str(existing.id), "updated": True}
        lead = Lead(tenant_id=leads_c.tenant_id, campaign_id=leads_c.id, phone_e164=phone, raw_phone=body.phone,
                    dedupe_key=phone, contact_name=body.name or None, business_name=body.business or None,
                    extra=_with_category({}, body.category), status="pending")
        session.add(lead)
        await session.flush()
        return {"id": str(lead.id), "updated": False}


@router.delete("/leads/{lead_id}", dependencies=[Depends(require_console)])
async def delete_lead(lead_id: uuid.UUID) -> dict:
    async with session_scope() as session:
        leads_c, _ = await _campaigns(session)
        lead = await session.get(Lead, lead_id)
        if lead is None or lead.campaign_id != leads_c.id:
            raise HTTPException(status_code=404, detail="Lead not found")
        if lead.attempts_count or lead.status not in ("pending",):
            # Called leads keep their history (and DNC must never be forgotten).
            raise HTTPException(status_code=409, detail="Only leads that were never called can be removed")
        await session.delete(lead)
    return {"deleted": True}


@router.post("/leads/upload", dependencies=[Depends(require_console)])
async def upload_leads(file: UploadFile) -> dict:
    name = (file.filename or "").lower()
    if not name.endswith((".xlsx", ".csv")):
        raise HTTPException(status_code=415, detail="Upload an .xlsx or .csv file")
    data = await file.read()
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="File too large (max 5 MB)")
    buffer = io.BytesIO(data)
    columns = list((pd.read_excel(buffer, nrows=0) if name.endswith(".xlsx") else pd.read_csv(buffer, nrows=0)).columns)
    mapping = guess_column_mapping(columns)
    if "phone" not in mapping:
        raise HTTPException(status_code=422,
                            detail=f"No phone column found. Columns in the file: {', '.join(map(str, columns))}. "
                                   "Name it 'Phone' (see the template).")
    parsed, report = import_leads(data, file.filename, mapping)
    for row in parsed:
        for key in [k for k in row.extra if str(k).strip().lower() in _CATEGORY_HEADERS]:
            row.extra = _with_category({k: v for k, v in row.extra.items() if k != key}, str(row.extra[key]))
    async with session_scope() as session:
        leads_c, _ = await _campaigns(session)
        result = await CampaignService(session).import_leads(leads_c.tenant_id, leads_c.id, parsed)
    return {
        "columns": mapping, "rows": report.total_rows, "valid": report.valid_rows,
        "duplicates_in_file": report.duplicate_within_file,
        "errors": [{"row": e.row_index, "reason": e.reason} for e in report.errors][:50],
        **result,
    }


@router.get("/leads/template.xlsx", dependencies=[Depends(require_console)])
async def lead_template() -> Response:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(["Name", "Phone", "Business", "Product category"])
    ws.append(["Naman", "+91 98765 43210", "Demo Business", "Perfumes"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7C3AED")
    for col, width in zip("ABCD", (22, 20, 26, 24)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="leads-template.xlsx"'})


# --------------------------------------------------------------------------- calls


class TestCall(BaseModel):
    phone: str
    name: str = ""
    business: str = ""
    category: str = ""


@router.post("/calls/test", dependencies=[Depends(require_console)])
async def test_call(body: TestCall) -> dict:
    """Call one number right now, outside the queue (any time of day)."""
    phone = _normalize_phone(body.phone)
    async with session_scope() as session:
        _, test_c = await _campaigns(session)
        if await CampaignService(session).leads.is_suppressed(test_c.tenant_id, phone):
            raise HTTPException(status_code=409, detail="This number asked not to be called (do-not-call list)")
        busy = (await session.execute(
            select(func.count()).select_from(CallAttempt)
            .where(CallAttempt.campaign_id == test_c.id, CallAttempt.status.in_(ACTIVE_CALL_STATUSES),
                   CallAttempt.created_at > datetime.fromtimestamp(time.time() - 600, timezone.utc))
        )).scalar_one()
        if busy:
            raise HTTPException(status_code=409, detail="A test call is already in progress — wait for it to end")
        for stale in (await session.execute(select(CallAttempt).where(
            CallAttempt.campaign_id == test_c.id, CallAttempt.status.in_(ACTIVE_CALL_STATUSES)
        ))).scalars():
            stale.status = "canceled"  # an earlier test whose end was never reported
        for other in (await session.execute(select(Lead).where(Lead.campaign_id == test_c.id,
                                                                Lead.status.in_(["pending", "queued"])))).scalars():
            other.status = "completed"  # exactly one dialable lead: this one
        lead = (await session.execute(
            select(Lead).where(Lead.campaign_id == test_c.id, Lead.dedupe_key == phone)
        )).scalar_one_or_none()
        if lead is None:
            lead = Lead(tenant_id=test_c.tenant_id, campaign_id=test_c.id, phone_e164=phone, raw_phone=body.phone,
                        dedupe_key=phone, status="pending")
            session.add(lead)
        lead.contact_name, lead.business_name, lead.status = body.name or None, body.business or None, "pending"
        lead.extra = _with_category({k: v for k, v in (lead.extra or {}).items() if k != "product_category"}, body.category)
        test_c.status = "active"
        await session.flush()
        try:
            dispatched = await _dispatch_one(session, test_c)
        finally:
            test_c.status = "paused"
        attempt = (await session.execute(
            select(CallAttempt).where(CallAttempt.lead_id == lead.id).order_by(CallAttempt.created_at.desc()).limit(1)
        )).scalar_one_or_none()
        if not dispatched:
            detail = attempt.error if attempt is not None and attempt.error else "The call could not be placed"
            raise HTTPException(status_code=502, detail=detail[:300])
        return {"attempt_id": str(attempt.id), "status": attempt.status}


@router.get("/calls", dependencies=[Depends(require_console)])
async def list_calls(limit: int = 100) -> list[dict]:
    rows = await _console_rows()
    return [_row_dict(r) for r in reversed(rows[-limit:])]


@router.get("/calls/{attempt_id}", dependencies=[Depends(require_console)])
async def call_detail(attempt_id: uuid.UUID) -> dict:
    rows = await _console_rows()
    row = next((r for r in rows if r.attempt_id == str(attempt_id)), None)
    if row is None:
        raise HTTPException(status_code=404, detail="Call not found")
    lines = []
    for line in row.transcript.splitlines():
        speaker, _, text = line.partition(": ")
        lines.append({"speaker": "agent" if speaker == "Agent" else "customer", "text": text})
    return {**_row_dict(row), "facts": row.facts, "transcript": lines}


async def _end_attempt(session, attempt: CallAttempt, reason: str) -> None:
    """Record the end of a call whose end the provider never reported to us (or that we
    ended ourselves), the same way the hangup webhook does — so the console stops
    showing it as live and the dialer moves on."""
    talked = (await session.execute(
        select(func.count()).select_from(Conversation).where(Conversation.call_attempt_id == attempt.id)
    )).scalar_one()
    status = "completed" if talked else ("canceled" if reason == "ended_from_console" else "no_answer")
    await CallRepository(session).mark_attempt_status(attempt.id, status, ended_at=datetime.now(timezone.utc))
    await finalize_call(session, attempt.id, call_status=status, ended_reason=reason)


@router.post("/calls/{attempt_id}/hangup", dependencies=[Depends(require_console)])
async def hang_up(attempt_id: uuid.UUID) -> dict:
    """End a call from the console (the End call button)."""
    async with session_scope() as session:
        leads_c, test_c = await _campaigns(session)
        attempt = await session.get(CallAttempt, attempt_id)
        if attempt is None or attempt.campaign_id not in (leads_c.id, test_c.id):
            raise HTTPException(status_code=404, detail="Call not found")
        if attempt.status not in ACTIVE_CALL_STATUSES:
            return {"ended": True}
        campaign = leads_c if attempt.campaign_id == leads_c.id else test_c
        if attempt.provider_call_id:
            try:
                await get_telephony_provider(campaign.telephony_provider).hangup_call(attempt.provider_call_id)
            except Exception as exc:  # noqa: BLE001
                logger.warning("console_hangup_failed", extra={"call_attempt_id": str(attempt.id), "error": str(exc)[:200]})
                raise HTTPException(status_code=502, detail=f"The phone provider refused to end the call: {exc}") from exc
    # The provider's hangup webhook normally lands within a second or two; record the
    # end ourselves if it doesn't, so the button always ends the call in the console.
    await asyncio.sleep(3)
    async with session_scope() as session:
        attempt = await session.get(CallAttempt, attempt_id)
        if attempt is not None and attempt.status in ACTIVE_CALL_STATUSES:
            await _end_attempt(session, attempt, "ended_from_console")
    return {"ended": True}


@router.get("/export.xlsx", dependencies=[Depends(require_console)])
async def export_excel() -> Response:
    """All leads (current status + latest result), every call, and the interested leads."""
    import tempfile
    from pathlib import Path

    from openpyxl import load_workbook

    rows = await _console_rows()
    leads = await list_leads()
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "calls.xlsx"
        write_workbook(rows, path)
        wb = load_workbook(path)
    ws = wb.create_sheet("All leads", 0)
    ws.append(["Name", "Business", "Phone", "Lead result", "Status", "Calls made", "Summary", "Added"])
    for lead in leads:
        ws.append([lead["name"], lead["business"], lead["phone"], lead["result"], lead["status"], lead["attempts"],
                   lead["summary"], lead["added"][:16].replace("T", " ")])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7C3AED")
    for col, width in zip("ABCDEFGH", (20, 22, 16, 20, 16, 11, 50, 17)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    stamp = datetime.now().strftime("%Y-%m-%d")
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="sales-agent-leads-{stamp}.xlsx"'})


# --------------------------------------------------------------------------- dialer


@router.post("/campaign/start", dependencies=[Depends(require_console)])
async def start_calling() -> dict:
    if not _public_base().startswith("https://"):
        raise HTTPException(status_code=503, detail="Server has no public https URL (PUBLIC_BASE_URL) — start the tunnel")
    async with session_scope() as session:
        leads_c, _ = await _campaigns(session)
        leads_c.status = "active"
    return {"status": "active"}


@router.post("/campaign/pause", dependencies=[Depends(require_console)])
async def pause_calling() -> dict:
    async with session_scope() as session:
        leads_c, _ = await _campaigns(session)
        leads_c.status = "paused"
    return {"status": "paused"}


class CallingHours(BaseModel):
    start: str
    end: str


@router.put("/campaign/hours", dependencies=[Depends(require_console)])
async def set_calling_hours(body: CallingHours) -> dict:
    for value in (body.start, body.end):
        try:
            datetime.strptime(value, "%H:%M")
        except ValueError:
            raise HTTPException(status_code=422, detail=f"Use HH:MM, got {value!r}") from None
    if body.start >= body.end:
        raise HTTPException(status_code=422, detail="Start must be before end")
    async with session_scope() as session:
        leads_c, _ = await _campaigns(session)
        leads_c.call_window_start, leads_c.call_window_end = body.start, body.end
    return {"start": body.start, "end": body.end}


@router.get("/system", dependencies=[Depends(require_console)])
async def system_status() -> dict:
    s = get_settings()
    out = {"telephony": s.telephony_provider, "ai_backend": s.ai_backend, "llm": s.openai_llm_model,
           "public_url_ok": _public_base().startswith("https://"), "balance": None}
    if s.telephony_provider == "vobiz":
        try:
            provider = get_telephony_provider("vobiz")
            out["balance"] = (await asyncio.wait_for(provider.get_balance("INR"), 8)).get("available_balance")
        except Exception:  # noqa: BLE001
            out["balance"] = None
    return out


async def run_console_dialer() -> None:
    """Calls the Leads campaign's pending leads one at a time while it is running and
    within its calling hours. `dispatch_next_call` holds every rule (call window,
    concurrency = 1, DNC suppression, idempotent attempts)."""
    logger.info("console_dialer_started")
    tick = 0
    while True:
        tick += 1
        try:
            if tick % RECONCILE_EVERY_TICKS == 0:
                await _reconcile_live_calls()
            async with session_scope() as session:
                leads_c, _ = await _campaigns(session)
                if leads_c.status == "active" and _public_base().startswith("https://"):
                    await _dispatch_one(session, leads_c)
        except asyncio.CancelledError:
            raise
        except Exception:  # noqa: BLE001 - one bad tick must not stop the dialer
            logger.exception("console_dialer_tick_failed")
        await asyncio.sleep(DIALER_POLL_SECONDS)


_not_live_checks: dict[uuid.UUID, int] = {}


async def _reconcile_live_calls() -> None:
    """Calls the console thinks are still live, checked against the provider. Seen on a
    real Vobiz call (2026-10-22): the customer hung up but Vobiz kept the call open and
    sent no hangup webhook, so the console showed it as live (and blocked the next test
    call) until it was ended by hand. Two "not live" answers in a row end it here; a
    call past its maximum length is hung up."""
    s = get_settings()
    now = datetime.now(timezone.utc)
    async with session_scope() as session:
        leads_c, test_c = await _campaigns(session)
        attempts = (await session.execute(
            select(CallAttempt).where(CallAttempt.campaign_id.in_([leads_c.id, test_c.id]),
                                      CallAttempt.status.in_(ACTIVE_CALL_STATUSES))
        )).scalars().all()
        for attempt in attempts:
            age = (now - attempt.created_at).total_seconds()
            if age < RECONCILE_AFTER_SECONDS:
                continue
            if not attempt.provider_call_id:
                if age > 120:  # never reached the provider
                    await _end_attempt(session, attempt, "never_placed")
                continue
            campaign = leads_c if attempt.campaign_id == leads_c.id else test_c
            provider = get_telephony_provider(campaign.telephony_provider)
            try:
                status = await asyncio.wait_for(provider.get_call_status(attempt.provider_call_id), 10)
            except Exception:  # noqa: BLE001 - provider can't tell (or no API): leave it to the webhook
                continue
            if status in (CallStatus.IN_PROGRESS, CallStatus.RINGING):
                _not_live_checks.pop(attempt.id, None)
                if age > s.default_max_call_duration_seconds + 60:
                    with contextlib.suppress(Exception):
                        await provider.hangup_call(attempt.provider_call_id)
                continue
            _not_live_checks[attempt.id] = _not_live_checks.get(attempt.id, 0) + 1
            if _not_live_checks[attempt.id] >= 2:
                _not_live_checks.pop(attempt.id, None)
                logger.warning("console_call_reconciled", extra={"call_attempt_id": str(attempt.id)})
                await _end_attempt(session, attempt, "provider_reports_ended")
