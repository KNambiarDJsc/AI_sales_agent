"""Excel call log: one row per call with what happened, plus a sheet of the leads who
agreed to talk to the sales team.

The workbook is rebuilt from the database (the source of truth) rather than appended
to, so it can't double-log a call that `finalize_call` finished twice (media teardown,
then the hangup webhook), and a file that was open in Excel — Windows locks it — just
catches up on the next rebuild. `schedule_call_log_refresh()` is called by
`finalize_call` and rebuilds a few seconds later, after the call's last turns have been
written (they are persisted fire-and-forget). `python scripts/export_call_log.py`
rebuilds it on demand.

"Agreed to sales team?" is decided by the application, not taken from the LLM: Yes if
the conversation reached one of the script's `sales_followup_states` (config) or a
recorded qualification says a sales follow-up is required; DNC and callbacks are
reported as such, and calls that never connected by the provider's status.
"""
from __future__ import annotations

import asyncio
import logging
import re
import threading
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from config.settings import get_settings
from database.models import CallAttempt, Callback, Campaign, Conversation, Lead, Qualification, TranscriptSegment, Turn
from database.session import session_scope
from orchestrator.state_machine import load_script_by_id

logger = logging.getLogger(__name__)

REFRESH_DELAY_SECONDS = 3.0
_write_lock = threading.Lock()
_pending: asyncio.Task | None = None

CALL_COLUMNS = [
    ("Call time", 19), ("Name", 18), ("Business", 20), ("Phone", 16), ("Agreed to sales team?", 20),
    ("Outcome", 24), ("Call status", 13), ("Talk time (s)", 12), ("Final state", 15), ("Path through script", 45),
    ("Callback time", 19), ("Summary", 40), ("Facts captured", 40), ("Transcript", 90), ("Campaign", 22),
    ("Provider", 11), ("Call attempt id", 38),
]
YES_COLUMNS = [("Call time", 19), ("Name", 18), ("Business", 20), ("Phone", 16), ("What they said", 70),
               ("Summary", 40), ("Call attempt id", 38)]
NEVER_CONNECTED = {"no_answer": "No answer", "busy": "Busy", "failed": "Call failed", "canceled": "Canceled"}


@dataclass
class CallRow:
    call_time: datetime
    name: str
    business: str
    phone: str
    agreed: str
    outcome: str
    call_status: str
    talk_seconds: int | None
    final_state: str
    path: str
    callback_time: str
    summary: str
    facts: str
    transcript: str
    last_customer_words: str
    campaign: str
    provider: str
    attempt_id: str


def _agreed(attempt: CallAttempt, path: list[str], followup_states: set[str], qualification, callback) -> str:
    if attempt.outcome == "do_not_call":
        return "Do not call"
    if followup_states & set(path) or (qualification is not None and qualification.sales_followup_required):
        return "Yes"
    if callback is not None:
        return "Callback requested"
    if not path:
        return NEVER_CONNECTED.get(attempt.status, "Not connected")
    return "No"


def _local(dt: datetime | None, tz: str) -> datetime | None:
    if dt is None:
        return None
    try:
        return dt.astimezone(ZoneInfo(tz)).replace(tzinfo=None)
    except Exception:  # noqa: BLE001 - a bad campaign timezone must not break the log
        return dt.replace(tzinfo=None)


def _is_logged(provider: str, campaign_name: str) -> bool:
    settings = get_settings()
    skip_providers = {p.strip() for p in settings.call_log_skip_providers.split(",") if p.strip()}
    if provider in skip_providers:
        return False
    pattern = settings.call_log_skip_campaigns
    return not (pattern and re.search(pattern, campaign_name or ""))


async def collect_rows() -> list[CallRow]:
    rows: list[CallRow] = []
    scripts: dict[str, set[str]] = {}
    async with session_scope() as session:
        attempts = (await session.execute(select(CallAttempt).order_by(CallAttempt.created_at))).scalars().all()
        for attempt in attempts:
            lead = await session.get(Lead, attempt.lead_id)
            campaign = await session.get(Campaign, attempt.campaign_id)
            if lead is None or campaign is None or not _is_logged(attempt.provider, campaign.name):
                continue
            if campaign.script_id not in scripts:
                try:
                    scripts[campaign.script_id] = set(load_script_by_id(campaign.script_id).sales_followup_states)
                except (FileNotFoundError, KeyError, ValueError):
                    scripts[campaign.script_id] = set()
            conversation = (await session.execute(
                select(Conversation).where(Conversation.call_attempt_id == attempt.id)
                .order_by(Conversation.created_at.desc()).limit(1)
            )).scalar_one_or_none()

            path: list[str] = []
            lines: list[str] = []
            last_customer = ""
            qualification = callback = None
            if conversation is not None:
                turns = (await session.execute(
                    select(Turn).where(Turn.conversation_id == conversation.id).order_by(Turn.turn_index)
                )).scalars().all()
                segments = {
                    s.turn_id: s.text for s in (await session.execute(
                        select(TranscriptSegment).where(TranscriptSegment.conversation_id == conversation.id)
                    )).scalars()
                }
                for turn in turns:
                    states = [turn.state]
                    if turn.speaker == "agent" and turn.raw_llm_output:
                        states.append(str(turn.raw_llm_output.get("state") or turn.state))
                    for state in states:
                        if not path or path[-1] != state:
                            path.append(state)
                    text = segments.get(turn.id, "")
                    if text:
                        lines.append(f"{'Agent' if turn.speaker == 'agent' else 'Customer'}: {text}")
                        if turn.speaker == "customer":
                            last_customer = text
                if conversation.current_state and (not path or path[-1] != conversation.current_state):
                    path.append(conversation.current_state)
                qualification = (await session.execute(
                    select(Qualification).where(Qualification.conversation_id == conversation.id)
                    .order_by(Qualification.created_at.desc()).limit(1)
                )).scalar_one_or_none()
                callback = (await session.execute(
                    select(Callback).where(Callback.conversation_id == conversation.id)
                    .order_by(Callback.created_at.desc()).limit(1)
                )).scalar_one_or_none()

            talk = None
            if conversation is not None and conversation.ended_at is not None:
                talk = int((conversation.ended_at - conversation.created_at).total_seconds())
            tz = campaign.timezone
            callback_time = _local(callback.requested_time, callback.timezone or tz) if callback else None
            rows.append(CallRow(
                call_time=_local(attempt.created_at, tz),
                name=lead.contact_name or "",
                business=lead.business_name or "",
                phone=lead.phone_e164,
                agreed=_agreed(attempt, path, scripts[campaign.script_id], qualification, callback),
                outcome=attempt.outcome or "",
                call_status=attempt.status,
                talk_seconds=talk,
                final_state=conversation.current_state if conversation else "",
                path=" → ".join(path),
                callback_time=callback_time.strftime("%Y-%m-%d %H:%M") if callback_time else "",
                summary=(qualification.summary or "") if qualification else "",
                facts=", ".join(f"{k}={v}" for k, v in (qualification.facts or {}).items()) if qualification else "",
                transcript="\n".join(lines),
                last_customer_words=last_customer,
                campaign=campaign.name,
                provider=attempt.provider,
                attempt_id=str(attempt.id),
            ))
    return rows


def write_workbook(rows: list[CallRow], path: Path) -> None:
    wb = Workbook()
    calls = wb.active
    calls.title = "Calls"
    yes = wb.create_sheet("Agreed to sales team")
    header_font, header_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    for sheet, columns in ((calls, CALL_COLUMNS), (yes, YES_COLUMNS)):
        sheet.append([name for name, _ in columns])
        for i, (_, width) in enumerate(columns, start=1):
            sheet.column_dimensions[get_column_letter(i)].width = width
            cell = sheet.cell(row=1, column=i)
            cell.font, cell.fill = header_font, header_fill
        sheet.freeze_panes = "A2"

    green, red = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFC7CE")
    for row in reversed(rows):  # newest first
        calls.append([
            row.call_time, row.name, row.business, row.phone, row.agreed, row.outcome, row.call_status,
            row.talk_seconds, row.final_state, row.path, row.callback_time, row.summary, row.facts, row.transcript,
            row.campaign, row.provider, row.attempt_id,
        ])
        r = calls.max_row
        calls.cell(row=r, column=1).number_format = "yyyy-mm-dd hh:mm"
        if row.agreed == "Yes":
            calls.cell(row=r, column=5).fill = green
            yes.append([row.call_time, row.name, row.business, row.phone, row.last_customer_words, row.summary,
                        row.attempt_id])
            yes.cell(row=yes.max_row, column=1).number_format = "yyyy-mm-dd hh:mm"
        elif row.agreed == "Do not call":
            calls.cell(row=r, column=5).fill = red
        for c in (10, 12, 13, 14):
            calls.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)  # atomic swap; raises PermissionError if Excel has the file open


async def refresh_call_log() -> Path | None:
    """Rebuild the workbook now. Never raises — the log must not affect a live call."""
    settings = get_settings()
    if not settings.call_log_enabled:
        return None
    path = Path(settings.call_log_path)
    try:
        rows = await collect_rows()

        def _write():
            with _write_lock:
                write_workbook(rows, path)

        await asyncio.get_running_loop().run_in_executor(None, _write)
        logger.info("call_log_written", extra={"rows": len(rows), "path": str(path)})
        return path
    except PermissionError:
        logger.warning("call_log_locked", extra={"path": str(path), "hint": "close the file in Excel; next call retries"})
    except Exception:  # noqa: BLE001
        logger.exception("call_log_failed")
    return None


def schedule_call_log_refresh(delay: float = REFRESH_DELAY_SECONDS) -> None:
    """Rebuild the log `delay` seconds from now; repeated calls within that window
    collapse into one rebuild. Safe to call from any coroutine; a no-op without a
    running event loop."""
    global _pending
    try:
        loop = asyncio.get_running_loop()
    except RuntimeError:
        return
    if _pending is not None and not _pending.done():
        _pending.cancel()

    async def _later():
        await asyncio.sleep(delay)
        await refresh_call_log()

    _pending = loop.create_task(_later())
